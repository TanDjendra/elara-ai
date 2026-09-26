import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml

doc = docx.Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Header / Kop Surat
header_p = doc.add_paragraph()
header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_org = header_p.add_run('ORGANISASI DEVELOPER INDONESIA\n')
run_org.bold = True
run_org.font.size = Pt(14)
run_org.font.name = 'Arial'
run_org.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

run_sub = header_p.add_run('Jl. Teknologi No. 1, Jakarta Selatan | Email: info@devorg.id | Web: www.devorg.id\n')
run_sub.font.size = Pt(9)
run_sub.font.name = 'Arial'
run_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

# Horizontal line under header
pBdr_xml = '<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="12" w:space="4" w:color="111827"/></w:pBdr>'
header_p._p.get_or_add_pPr().append(parse_xml(pBdr_xml))

doc.add_paragraph()

# Metadata
meta_p = doc.add_paragraph()
meta_p.paragraph_format.space_after = Pt(12)

def add_meta(p, label, val):
    r1 = p.add_run(f'{label:<12}: ')
    r1.font.name = 'Arial'
    r1.font.size = Pt(11)
    r2 = p.add_run(f'{val}\n')
    r2.font.name = 'Arial'
    r2.font.size = Pt(11)

add_meta(meta_p, 'Nomor', '001/DEV-ORG/UND/III/2025')
add_meta(meta_p, 'Lampiran', '-')
add_meta(meta_p, 'Hal', 'Undangan Acara Perkumpulan Developer')

# Target
to_p = doc.add_paragraph()
to_p.paragraph_format.space_after = Pt(12)
r_to = to_p.add_run('Kepada Yth.\nBapak/Ibu / Anggota Developer\ndi Tempat\n')
r_to.font.name = 'Arial'
r_to.font.size = Pt(11)

# Opening
body1 = doc.add_paragraph()
body1.paragraph_format.space_after = Pt(8)
body1.paragraph_format.line_spacing = 1.15
r_b1 = body1.add_run('Dengan hormat,\n\nSehubungan dengan upaya mempererat tali silaturahmi serta berbagi wawasan mengenai perkembangan teknologi terkini, kami dari pengurus Organisasi Developer Indonesia bermaksud mengundang Bapak/Ibu/Rekan-rekan sekalian untuk hadir dalam acara Perkumpulan Developer.')
r_b1.font.name = 'Arial'
r_b1.font.size = Pt(11)

# Event details table
table = doc.add_table(rows=4, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

details = [
    ('Nama Acara', 'Perkumpulan Developer'),
    ('Hari, Tanggal', 'Rabu, 5 Maret 2025'),
    ('Waktu', '19.00 WIB - Selesai'),
    ('Tempat / Media', 'Ruang Rapat Utama / Discord Community Server')
]

for i, (label, val) in enumerate(details):
    row = table.rows[i]
    cell_lbl = row.cells[0]
    cell_lbl.width = Inches(1.8)
    p_lbl = cell_lbl.paragraphs[0]
    r_l = p_lbl.add_run(label)
    r_l.bold = True
    r_l.font.name = 'Arial'
    r_l.font.size = Pt(10.5)
    r_l.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    
    cell_val = row.cells[1]
    cell_val.width = Inches(4.2)
    p_val = cell_val.paragraphs[0]
    r_v = p_val.add_run(val)
    r_v.font.name = 'Arial'
    r_v.font.size = Pt(10.5)

doc.add_paragraph()

# Closing
body2 = doc.add_paragraph()
body2.paragraph_format.space_after = Pt(24)
body2.paragraph_format.line_spacing = 1.15
r_b2 = body2.add_run('Mengingat pentingnya acara ini untuk kemajuan ekosistem dan kolaborasi developer kita, kehadiran Rekan-rekan sekalian sangat kami harapkan.\n\nDemikian surat undangan ini kami sampaikan. Atas perhatian dan kehadirannya, kami ucapkan terima kasih.')
r_b2.font.name = 'Arial'
r_b2.font.size = Pt(11)

# Signatures
sig_table = doc.add_table(rows=1, cols=2)
sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cell_left = sig_table.rows[0].cells[0]
cell_right = sig_table.rows[0].cells[1]
cell_left.width = Inches(3.0)
cell_right.width = Inches(3.0)

p_r = cell_right.paragraphs[0]
p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sig = p_r.add_run('Hormat kami,\nPanitia Pelaksana\n\n\n\n\n')
r_sig.font.name = 'Arial'
r_sig.font.size = Pt(11)

r_name = p_r.add_run('( Tan )\n')
r_name.bold = True
r_name.font.name = 'Arial'
r_name.font.size = Pt(11)

r_title = p_r.add_run('Ketua Panitia')
r_title.font.name = 'Arial'
r_title.font.size = Pt(10)

doc.save('Surat_Undangan_Perkumpulan_Developer.docx')
print('SUCCESS')
