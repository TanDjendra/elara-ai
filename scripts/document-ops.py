"""Bounded local DOCX/PDF/XLSX inspection and precise edits for ELARA.

Run through DSH's existing approved shell tool. This script never sends data,
overwrites the input, or changes the DSH authorization boundary.
"""

from __future__ import annotations

import argparse
import csv
import datetime
import hashlib
import json
import os
import pathlib
import re
import sys
import tempfile
import zipfile
from decimal import Decimal, InvalidOperation
from typing import Iterable


MAX_INPUT_BYTES = 25 * 1024 * 1024
MAX_TEXT_CHARS = 30_000
MAX_XLSX_SHEETS = 20
MAX_XLSX_ROWS = 200
MAX_XLSX_COLUMNS = 50
MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_XLSX_CREATE_COLUMNS = 40
MAX_XLSX_CREATE_ROWS = 10_000
MAX_XLSX_SPEC_BYTES = 1024 * 1024
XLSX_COLUMN_TYPES = {"text", "integer", "number", "currency", "percent", "date", "boolean"}


def safe_input(raw: str) -> pathlib.Path:
    source = pathlib.Path(raw).resolve(strict=True)
    if not source.is_file() or source.stat().st_size > MAX_INPUT_BYTES:
        raise ValueError("DOCUMENT_INVALID")
    if source.suffix.lower() not in {".docx", ".pdf", ".xlsx"}:
        raise ValueError("DOCUMENT_FORMAT_UNSUPPORTED")
    return source


def safe_output(raw: str, source: pathlib.Path) -> pathlib.Path:
    target = pathlib.Path(raw).resolve(strict=False)
    if target == source or target.exists() or target.suffix.lower() != source.suffix.lower():
        raise ValueError("DOCUMENT_OUTPUT_INVALID")
    if not target.parent.is_dir():
        raise ValueError("DOCUMENT_OUTPUT_DIRECTORY_MISSING")
    return target


def new_xlsx_output(raw: str) -> pathlib.Path:
    target = pathlib.Path(raw).resolve(strict=False)
    if target.suffix.lower() != ".xlsx" or target.exists():
        raise ValueError("DOCUMENT_OUTPUT_INVALID")
    if not target.parent.is_dir():
        raise ValueError("DOCUMENT_OUTPUT_DIRECTORY_MISSING")
    return target


def short_text(value: object, maximum: int, code: str) -> str:
    if (not isinstance(value, str) or not value.strip() or len(value) > maximum
            or any(ord(character) < 32 for character in value)):
        raise ValueError(code)
    return value.strip()


def creation_spec(args: argparse.Namespace) -> dict:
    if args.spec:
        source = pathlib.Path(args.spec).resolve(strict=True)
        if source.suffix.lower() != ".json" or not source.is_file() or source.stat().st_size > MAX_XLSX_SPEC_BYTES:
            raise ValueError("DOCUMENT_SPEC_INVALID")
        try:
            spec = json.loads(source.read_text(encoding="utf-8"))
        except (UnicodeError, json.JSONDecodeError) as error:
            raise ValueError("DOCUMENT_SPEC_INVALID") from error
        if not isinstance(spec, dict) or set(spec) - {"title", "sheet", "subtitle", "columns", "rows"}:
            raise ValueError("DOCUMENT_SPEC_INVALID")
        return spec
    try:
        labels = next(csv.reader([args.columns]))
    except csv.Error as error:
        raise ValueError("DOCUMENT_COLUMNS_INVALID") from error
    return {"title": args.title or "Data", "sheet": args.sheet or "Data",
            "columns": [{"key": f"column_{index}", "label": label, "type": "text"}
                        for index, label in enumerate(labels, start=1)], "rows": []}


def normalized_creation_spec(raw: dict) -> tuple[str, str, str, list[dict], list[dict]]:
    title = short_text(raw.get("title", "Data"), 120, "DOCUMENT_TITLE_INVALID")
    sheet = short_text(raw.get("sheet", "Data"), 31, "DOCUMENT_SHEET_INVALID")
    if re.search(r"[\\/*?:\[\]]", sheet):
        raise ValueError("DOCUMENT_SHEET_INVALID")
    subtitle = raw.get("subtitle", "")
    if not isinstance(subtitle, str) or len(subtitle) > 180:
        raise ValueError("DOCUMENT_SUBTITLE_INVALID")
    columns = raw.get("columns")
    if not isinstance(columns, list) or not 1 <= len(columns) <= MAX_XLSX_CREATE_COLUMNS:
        raise ValueError("DOCUMENT_COLUMNS_INVALID")
    normalized: list[dict] = []
    keys: set[str] = set()
    for entry in columns:
        if not isinstance(entry, dict) or set(entry) - {"key", "label", "type", "symbol"}:
            raise ValueError("DOCUMENT_COLUMNS_INVALID")
        key = short_text(entry.get("key"), 60, "DOCUMENT_COLUMNS_INVALID")
        label = short_text(entry.get("label"), 100, "DOCUMENT_COLUMNS_INVALID")
        value_type = entry.get("type", "text")
        symbol = entry.get("symbol", "Rp")
        if key in keys or not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", key) or value_type not in XLSX_COLUMN_TYPES:
            raise ValueError("DOCUMENT_COLUMNS_INVALID")
        if ("symbol" in entry and value_type != "currency") or not isinstance(symbol, str) or not re.fullmatch(r"[A-Za-z$€£¥]{1,5}", symbol):
            raise ValueError("DOCUMENT_COLUMNS_INVALID")
        keys.add(key)
        normalized.append({"key": key, "label": label, "type": value_type, "symbol": symbol})
    rows = raw.get("rows", [])
    if not isinstance(rows, list) or len(rows) > MAX_XLSX_CREATE_ROWS:
        raise ValueError("DOCUMENT_ROWS_INVALID")
    for row in rows:
        if not isinstance(row, dict) or set(row) - keys:
            raise ValueError("DOCUMENT_ROWS_INVALID")
    return title, sheet, subtitle, normalized, rows


def typed_spreadsheet_value(value: object, value_type: str) -> object:
    if value is None or value == "":
        return None
    if value_type == "text":
        if not isinstance(value, (str, int, float, bool)) or len(str(value)) > 5_000:
            raise ValueError("DOCUMENT_VALUE_INVALID")
        return str(value)
    if value_type == "boolean":
        if isinstance(value, bool):
            return value
        if isinstance(value, str) and value.lower() in {"true", "false"}:
            return value.lower() == "true"
        raise ValueError("DOCUMENT_VALUE_INVALID")
    if value_type == "date":
        try:
            return datetime.date.fromisoformat(value)
        except (TypeError, ValueError) as error:
            raise ValueError("DOCUMENT_VALUE_INVALID") from error
    if isinstance(value, bool):
        raise ValueError("DOCUMENT_VALUE_INVALID")
    try:
        number = Decimal(str(value))
        if not number.is_finite() or abs(number) > Decimal("1e15"):
            raise InvalidOperation
        if value_type == "integer":
            if number != int(number):
                raise InvalidOperation
            return int(number)
        return float(number)
    except (InvalidOperation, ValueError, OverflowError) as error:
        raise ValueError("DOCUMENT_VALUE_INVALID") from error


def create_xlsx(target: pathlib.Path, raw_spec: dict) -> tuple[int, int]:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    title, sheet_name, subtitle, columns, rows = normalized_creation_spec(raw_spec)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.sheet_properties.tabColor = "24527A"
    sheet.sheet_view.showGridLines = False
    end_column = get_column_letter(len(columns))
    if len(columns) > 1:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
    for column_number in range(1, len(columns) + 1):
        sheet.cell(1, column_number).fill = PatternFill("solid", fgColor="14283F")
        sheet.cell(2, column_number).fill = PatternFill("solid", fgColor="EAF0F6")
    title_cell = sheet.cell(1, 1, title)
    title_cell.data_type = "s"
    title_cell.font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
    title_cell.alignment = Alignment(vertical="center", indent=1)
    sheet.row_dimensions[1].height = 34
    subtitle_cell = sheet.cell(2, 1, subtitle or "Data siap diisi")
    subtitle_cell.data_type = "s"
    subtitle_cell.font = Font(name="Aptos", size=10, color="50657B")
    subtitle_cell.alignment = Alignment(vertical="center", indent=2)
    sheet.row_dimensions[2].height = 24
    edge = Side(style="hair", color="E3EAF1")
    for column_number, column in enumerate(columns, start=1):
        heading = sheet.cell(3, column_number, column["label"])
        heading.data_type = "s"
        heading.fill = PatternFill("solid", fgColor="24527A")
        heading.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        heading.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        heading.border = Border(bottom=Side(style="medium", color="183A58"))
        samples = [str(row.get(column["key"], "")) for row in rows[:100]]
        width = min(38, max(13, len(column["label"]) + 4,
                             min(38, max((len(item) for item in samples), default=0) + 3)))
        sheet.column_dimensions[get_column_letter(column_number)].width = width
    sheet.row_dimensions[3].height = 30
    body_count = max(len(rows), 12)
    for index in range(body_count):
        row = rows[index] if index < len(rows) else {}
        excel_row = index + 4
        sheet.row_dimensions[excel_row].height = 22
        background = "FFFFFF" if index % 2 == 0 else "F5F8FC"
        for column_number, column in enumerate(columns, start=1):
            value = typed_spreadsheet_value(row.get(column["key"]), column["type"])
            cell = sheet.cell(excel_row, column_number, value)
            if column["type"] == "text" and value is not None:
                cell.data_type = "s"
            cell.fill = PatternFill("solid", fgColor=background)
            cell.font = Font(name="Aptos", size=10, color="223044")
            cell.border = Border(bottom=edge)
            cell.alignment = Alignment(vertical="center", wrap_text=column["type"] == "text",
                                       horizontal="right" if column["type"] in {"integer", "number", "currency", "percent"} else "left")
            cell.number_format = {
                "integer": "#,##0", "number": "#,##0.00", "currency": '#,##0.00',
                "percent": "0.0%", "date": "yyyy-mm-dd",
            }.get(column["type"], "General")
            if column["type"] == "currency":
                cell.number_format = f'"{column["symbol"]}" #,##0.00'
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:{end_column}{body_count + 3}"
    sheet.print_title_rows = "1:3"
    sheet.print_area = f"A1:{end_column}{body_count + 3}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.orientation = "landscape" if len(columns) > 6 else "portrait"
    temp_name = None
    try:
        with tempfile.NamedTemporaryFile(dir=target.parent, prefix=".elara-xlsx-", suffix=".xlsx",
                                         delete=False) as temporary:
            temp_name = pathlib.Path(temporary.name)
        workbook.save(temp_name)
        if temp_name.stat().st_size > MAX_INPUT_BYTES:
            raise ValueError("DOCUMENT_TOO_COMPLEX")
        check_xlsx_archive(temp_name)
        verified = load_workbook(temp_name, read_only=True, data_only=False)
        try:
            check = verified[sheet_name]
            if [check.cell(3, number).value for number in range(1, len(columns) + 1)] != [
                column["label"] for column in columns
            ] or check["A1"].value != title:
                raise ValueError("DOCUMENT_EDIT_UNVERIFIED")
        finally:
            verified.close()
        try:
            os.link(temp_name, target)
        except FileExistsError as error:
            raise ValueError("DOCUMENT_OUTPUT_INVALID") from error
    finally:
        workbook.close()
        if temp_name is not None:
            temp_name.unlink(missing_ok=True)
    return len(columns), len(rows)


def all_paragraphs(document: object) -> Iterable[object]:
    def walk(container: object) -> Iterable[object]:
        yield from container.paragraphs
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)

    yield from walk(document)
    for section in document.sections:
        yield from walk(section.header)
        yield from walk(section.footer)


def check_xlsx_archive(source: pathlib.Path, editing: bool = False) -> None:
    try:
        with zipfile.ZipFile(source) as archive:
            members = archive.infolist()
            if len(members) > 5_000 or sum(item.file_size for item in members) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValueError("DOCUMENT_TOO_COMPLEX")
            if editing and any(item.filename.lower().startswith((
                "xl/drawings/", "xl/charts/", "xl/pivottables/", "xl/slicer",
                "xl/embeddings/", "xl/externallinks/", "xl/vbaproject",
            )) for item in members):
                raise ValueError("DOCUMENT_COMPLEX_EDIT_UNSUPPORTED")
    except zipfile.BadZipFile as error:
        raise ValueError("DOCUMENT_INVALID") from error


def xlsx_value(value: object) -> str:
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value).replace("\r", " ").replace("\n", " ⏎ ")


def read_xlsx(source: pathlib.Path) -> str:
    from openpyxl import load_workbook

    check_xlsx_archive(source)
    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    lines: list[str] = []
    length = 0
    partial = len(workbook.worksheets) > MAX_XLSX_SHEETS
    try:
        for sheet in workbook.worksheets[:MAX_XLSX_SHEETS]:
            heading = f"[Sheet: {sheet.title}]"
            lines.append(heading)
            length += len(heading) + 1
            last_row = sheet.max_row or 0
            last_column = sheet.max_column or 0
            partial |= last_row > MAX_XLSX_ROWS or last_column > MAX_XLSX_COLUMNS
            if last_row < 1 or last_column < 1:
                continue
            for row in sheet.iter_rows(min_row=1, max_row=min(last_row, MAX_XLSX_ROWS),
                                       max_col=min(last_column, MAX_XLSX_COLUMNS)):
                for cell in row:
                    if cell.value is None:
                        continue
                    value = xlsx_value(cell.value)
                    if cell.data_type == "f":
                        value += " [rumus; hasil belum dihitung ulang]"
                    line = f"{cell.coordinate}: {value}"
                    lines.append(line)
                    length += len(line) + 1
                    if length >= MAX_TEXT_CHARS:
                        partial = True
                        break
                if length >= MAX_TEXT_CHARS:
                    break
            if length >= MAX_TEXT_CHARS:
                break
    finally:
        workbook.close()
    output = "\n".join(lines).strip()
    return output[:MAX_TEXT_CHARS] + (
        "\n[Cuplikan spreadsheet terpotong. Jangan menganggap seluruh isi sudah terbaca.]" if partial else ""
    )


def formula_signature(workbook: object) -> str:
    digest = hashlib.sha256()
    for sheet in workbook.worksheets:
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    digest.update(f"{sheet.title}!{cell.coordinate}:{cell.value}\n".encode("utf-8"))
    return digest.hexdigest()


def check_editable_xlsx(workbook: object) -> None:
    if len(workbook.worksheets) > 50 or sum(
        sheet.max_row * sheet.max_column for sheet in workbook.worksheets
    ) > 250_000:
        raise ValueError("DOCUMENT_TOO_COMPLEX")


def replace_xlsx(source: pathlib.Path, target: pathlib.Path, old: str, new: str) -> int:
    from openpyxl import load_workbook

    check_xlsx_archive(source, editing=True)
    workbook = load_workbook(source, data_only=False, keep_links=True)
    check_editable_xlsx(workbook)
    formulas = formula_signature(workbook)
    changed: dict[tuple[str, str], str] = {}
    count = 0
    for sheet in workbook.worksheets:
        for row in sheet:
            for cell in row:
                if cell.data_type == "f" or not isinstance(cell.value, str) or old not in cell.value:
                    continue
                count += cell.value.count(old)
                cell.value = cell.value.replace(old, new)
                cell.data_type = "s"
                changed[(sheet.title, cell.coordinate)] = cell.value
    if not count:
        raise ValueError("DOCUMENT_TEXT_NOT_FOUND")
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(target)
    reopened = load_workbook(target, data_only=False, keep_links=True)
    try:
        if (reopened.sheetnames != workbook.sheetnames or formula_signature(reopened) != formulas
                or any(reopened[sheet][coordinate].value != value
                       for (sheet, coordinate), value in changed.items())):
            raise ValueError("DOCUMENT_EDIT_UNVERIFIED")
    except Exception:
        target.unlink(missing_ok=True)
        raise
    finally:
        reopened.close()
    workbook.close()
    return count


def set_xlsx_cell(source: pathlib.Path, target: pathlib.Path, sheet_name: str,
                  coordinate: str, raw_value: str, value_type: str) -> None:
    from decimal import Decimal, InvalidOperation
    from openpyxl import load_workbook
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

    check_xlsx_archive(source, editing=True)
    if not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]{0,6}", coordinate):
        raise ValueError("DOCUMENT_CELL_INVALID")
    column, row = coordinate_from_string(coordinate.upper())
    if column_index_from_string(column) > 16_384 or row > 1_048_576:
        raise ValueError("DOCUMENT_CELL_INVALID")
    workbook = load_workbook(source, data_only=False, keep_links=True)
    check_editable_xlsx(workbook)
    if sheet_name not in workbook.sheetnames:
        raise ValueError("DOCUMENT_SHEET_NOT_FOUND")
    sheet = workbook[sheet_name]
    cell = sheet[coordinate.upper()]
    if cell.data_type == "f":
        raise ValueError("DOCUMENT_FORMULA_CELL")
    formulas = formula_signature(workbook)
    if value_type == "number":
        try:
            parsed = Decimal(raw_value)
            if not parsed.is_finite():
                raise InvalidOperation
            value = int(parsed) if parsed == int(parsed) else float(parsed)
        except (InvalidOperation, ValueError, OverflowError) as error:
            raise ValueError("DOCUMENT_VALUE_INVALID") from error
    elif value_type == "boolean":
        if raw_value.lower() not in {"true", "false"}:
            raise ValueError("DOCUMENT_VALUE_INVALID")
        value = raw_value.lower() == "true"
    else:
        value = raw_value
    cell.value = value
    if value_type == "text":
        cell.data_type = "s"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(target)
    reopened = load_workbook(target, data_only=False, keep_links=True)
    try:
        if (reopened.sheetnames != workbook.sheetnames or formula_signature(reopened) != formulas
                or reopened[sheet_name][coordinate.upper()].value != value):
            raise ValueError("DOCUMENT_EDIT_UNVERIFIED")
    except Exception:
        target.unlink(missing_ok=True)
        raise
    finally:
        reopened.close()
        workbook.close()


def replace_in_paragraph(paragraph: object, old: str, new: str) -> int:
    changes = 0
    search_end = None
    while True:
        runs = list(paragraph.runs)
        whole = "".join(run.text for run in runs)
        start = whole.rfind(old, 0, len(whole) if search_end is None else search_end)
        if start < 0:
            return changes
        end = start + len(old)
        positions = [(index, offset) for index, run in enumerate(runs)
                     for offset in range(len(run.text))]
        first_index, first_offset = positions[start]
        last_index, last_offset = positions[end - 1]
        first = runs[first_index]
        last = runs[last_index]
        prefix = first.text[:first_offset]
        suffix = last.text[last_offset + 1:]
        if first_index == last_index:
            first.text = prefix + new + suffix
        else:
            first.text = prefix + new
            for index in range(first_index + 1, last_index):
                runs[index].text = ""
            last.text = suffix
        changes += 1
        search_end = start


def read_document(source: pathlib.Path) -> str:
    if source.suffix.lower() == ".docx":
        from docx import Document

        document = Document(source)
        text = "\n".join(paragraph.text for paragraph in all_paragraphs(document))
    elif source.suffix.lower() == ".xlsx":
        return read_xlsx(source)
    else:
        import pymupdf

        with pymupdf.open(source) as document:
            if document.needs_pass:
                raise ValueError("DOCUMENT_PASSWORD_REQUIRED")
            text = "\n".join(document[index].get_text("text")
                             for index in range(min(document.page_count, 50)))
    return text[:MAX_TEXT_CHARS]


def replace_docx(source: pathlib.Path, target: pathlib.Path, old: str, new: str) -> int:
    from docx import Document

    document = Document(source)
    changes = sum(replace_in_paragraph(paragraph, old, new)
                  for paragraph in all_paragraphs(document))
    if not changes:
        raise ValueError("DOCUMENT_TEXT_NOT_FOUND")
    document.save(target)
    reopened = Document(target)
    content = "\n".join(paragraph.text for paragraph in all_paragraphs(reopened))
    if (old not in new and old in content) or new not in content:
        target.unlink(missing_ok=True)
        raise ValueError("DOCUMENT_EDIT_UNVERIFIED")
    return changes


def replace_pdf(source: pathlib.Path, target: pathlib.Path, old: str, new: str) -> int:
    import pymupdf

    with pymupdf.open(source) as document:
        if document.needs_pass:
            raise ValueError("DOCUMENT_PASSWORD_REQUIRED")
        changes = 0
        original_pages = document.page_count
        for page in document:
            locations = page.search_for(old)
            for rectangle in locations:
                page.add_redact_annot(rectangle, text=new, fontsize=9, fill=(1, 1, 1))
            if locations:
                page.apply_redactions(images=0, graphics=0)
                changes += len(locations)
        if not changes:
            raise ValueError("DOCUMENT_TEXT_NOT_FOUND")
        document.save(target, garbage=4, deflate=True)
    with pymupdf.open(target) as reopened:
        content = "\n".join(page.get_text("text") for page in reopened)
        if reopened.page_count != original_pages or (old not in new and old in content) or new not in content:
            target.unlink(missing_ok=True)
            raise ValueError("DOCUMENT_EDIT_UNVERIFIED")
    return changes


def main() -> int:
    parser = argparse.ArgumentParser()
    subcommands = parser.add_subparsers(dest="action", required=True)
    read = subcommands.add_parser("read")
    read.add_argument("--input", required=True)
    replace = subcommands.add_parser("replace")
    replace.add_argument("--input", required=True)
    replace.add_argument("--output", required=True)
    replace.add_argument("--old", required=True)
    replace.add_argument("--new", required=True)
    set_cell = subcommands.add_parser("set-cell")
    set_cell.add_argument("--input", required=True)
    set_cell.add_argument("--output", required=True)
    set_cell.add_argument("--sheet", required=True)
    set_cell.add_argument("--cell", required=True)
    set_cell.add_argument("--value", required=True)
    set_cell.add_argument("--type", choices=["text", "number", "boolean"], default="text")
    create = subcommands.add_parser("create-xlsx")
    create.add_argument("--output", required=True)
    create_source = create.add_mutually_exclusive_group(required=True)
    create_source.add_argument("--columns")
    create_source.add_argument("--spec")
    create.add_argument("--title")
    create.add_argument("--sheet")
    args = parser.parse_args()
    try:
        if args.action == "create-xlsx":
            target = new_xlsx_output(args.output)
            columns, rows = create_xlsx(target, creation_spec(args))
            print(json.dumps({"outcome": "completed", "format": "xlsx",
                              "columns": columns, "rows": rows}))
            return 0
        source = safe_input(args.input)
        if args.action == "read":
            print(read_document(source))
            return 0
        if args.action == "set-cell":
            if source.suffix.lower() != ".xlsx":
                raise ValueError("DOCUMENT_FORMAT_UNSUPPORTED")
            target = safe_output(args.output, source)
            set_xlsx_cell(source, target, args.sheet, args.cell, args.value, args.type)
            print(json.dumps({"outcome": "completed", "format": "xlsx", "changes": 1}))
            return 0
        if not args.old or len(args.old) > 2_000 or len(args.new) > 2_000:
            raise ValueError("DOCUMENT_EDIT_INVALID")
        target = safe_output(args.output, source)
        if source.suffix.lower() == ".docx":
            count = replace_docx(source, target, args.old, args.new)
        elif source.suffix.lower() == ".pdf":
            count = replace_pdf(source, target, args.old, args.new)
        else:
            count = replace_xlsx(source, target, args.old, args.new)
        print(json.dumps({"outcome": "completed", "format": source.suffix.lower()[1:],
                          "replacements": count}))
        return 0
    except Exception as error:
        # No source path, document text, or exception message is printed.
        code = str(error) if isinstance(error, ValueError) and str(error).startswith("DOCUMENT_") else "DOCUMENT_OPERATION_FAILED"
        print(code, file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
